"""
03_report.py
Restitution : graphiques (PNG) et classeur Excel des 5 modèles PD PME France.
"""
import json
import sys

import numpy as np
import pandas as pd
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from sklearn.metrics import roc_curve, roc_auc_score

OUT = sys.argv[1]
H = [1, 2, 3, 4, 5]
SERIES = ["#2a78d6", "#eb6834", "#1baf7a", "#eda100", "#e87ba4"]   # palette catégorielle validée, ordre fixe
INK, INK2, GRID = "#0b0b0b", "#52514e", "#e4e3df"

plt.rcParams.update({"font.size": 10, "axes.edgecolor": GRID, "axes.labelcolor": INK2,
                     "xtick.color": INK2, "ytick.color": INK2, "axes.grid": True,
                     "grid.color": GRID, "grid.linewidth": 0.6, "axes.spines.top": False,
                     "axes.spines.right": False, "figure.dpi": 150})

perf = pd.read_csv(f"{OUT}/performance.csv")
coefs = pd.read_csv(f"{OUT}/coefficients.csv")
iv = pd.read_csv(f"{OUT}/information_value.csv", index_col=0)
dec = pd.read_csv(f"{OUT}/calibration_deciles.csv")
grades = pd.read_csv(f"{OUT}/echelle_notation.csv")
imp = pd.read_csv(f"{OUT}/lightgbm_importance.csv", index_col=0)
woe = pd.read_csv(f"{OUT}/grilles_woe.csv")
P = np.load(f"{OUT}/pred_test.npy")
ids = pd.read_csv(f"{OUT}/test_ids.csv")
summ = json.load(open(f"{OUT}/summary.json"))

# 1. Courbes ROC (test)
fig, ax = plt.subplots(figsize=(6, 5))
for i, h in enumerate(H):
    y = ids[f"def_{h}y"]
    m = y.notna().values
    fpr, tpr, _ = roc_curve(y[m], P[m, i])
    g = 2 * roc_auc_score(y[m], P[m, i]) - 1
    ax.plot(fpr, tpr, color=SERIES[i], lw=2, label=f"PD {h} an{'s' if h > 1 else ''} - Gini {g:.1%}")
ax.plot([0, 1], [0, 1], color=INK2, lw=1, ls="--")
ax.set_xlabel("Taux de faux positifs"); ax.set_ylabel("Taux de vrais positifs")
ax.set_title("Pouvoir discriminant - échantillon de test", color=INK, loc="left")
ax.legend(frameon=False, loc="lower right")
fig.tight_layout(); fig.savefig(f"{OUT}/fig_roc.png"); plt.close(fig)

# 2. Calibration par décile (test), petits multiples
fig, axes = plt.subplots(1, 5, figsize=(15, 3.4), sharey=False)
for i, h in enumerate(H):
    d = dec[dec.horizon == f"{h}Y"]
    ax = axes[i]
    ax.bar(d.decile - 0.2, d.pd_moyenne * 100, width=0.4, color=SERIES[0], label="PD moyenne estimée")
    ax.bar(d.decile + 0.2, d.taux_observe * 100, width=0.4, color=SERIES[1], label="Taux de défaut observé")
    ax.set_title(f"PD {h} an{'s' if h > 1 else ''}", color=INK, loc="left")
    ax.set_xticks(range(1, 11)); ax.set_xlabel("Décile de risque")
    if i == 0:
        ax.set_ylabel("%")
axes[0].legend(frameon=False, fontsize=8, loc="upper left")
fig.suptitle("Calibration par décile - échantillon de test", x=0.01, ha="left", color=INK)
fig.tight_layout(); fig.savefig(f"{OUT}/fig_calibration.png"); plt.close(fig)

# 3. Structure par terme des PD par classe de notation
fig, ax = plt.subplots(figsize=(7, 5))
cmap = plt.get_cmap("Blues")
sel = grades[grades.n >= 50]
for k, (_, r) in enumerate(sel.iterrows()):
    col = cmap(0.3 + 0.7 * k / max(len(sel) - 1, 1))
    ax.plot(H, [r[f"pd_{h}y_moy"] * 100 for h in H], color=col, lw=2, marker="o", ms=5)
    ax.text(5.08, r["pd_5y_moy"] * 100, r["grade"], va="center", fontsize=8, color=INK2)
ax.set_yscale("log"); ax.set_xticks(H)
ax.set_xlabel("Horizon (années)"); ax.set_ylabel("PD cumulée moyenne (%, échelle log)")
ax.set_title("Structure par terme des PD cumulées par classe", color=INK, loc="left")
fig.tight_layout(); fig.savefig(f"{OUT}/fig_term_structure.png"); plt.close(fig)

# ------------------------------------------------------------------ Excel
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = Workbook()
HDR = PatternFill("solid", fgColor="1F3A5F")
HFONT = Font(bold=True, color="FFFFFF")
thin = Side(style="thin", color="D9D9D9")


def write_df(ws, df, r0=1, c0=1, pct=(), dec4=(), dec2=()):
    for j, col in enumerate(df.columns):
        c = ws.cell(row=r0, column=c0 + j, value=str(col))
        c.fill, c.font = HDR, HFONT
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for i, row in enumerate(df.itertuples(index=False), start=1):
        for j, v in enumerate(row):
            if isinstance(v, (np.floating, float)) and np.isnan(v):
                v = None
            elif isinstance(v, np.generic):
                v = v.item()
            c = ws.cell(row=r0 + i, column=c0 + j, value=v)
            c.border = Border(bottom=thin)
            name = df.columns[j]
            if name in pct:
                c.number_format = "0.00%"
            elif name in dec4:
                c.number_format = "0.0000"
            elif name in dec2:
                c.number_format = "0.00"
    for j, col in enumerate(df.columns):
        width = max(len(str(col)), *(len(str(x)) for x in df.iloc[:, j].head(200))) if len(df) else 12
        ws.column_dimensions[get_column_letter(c0 + j)].width = min(max(10, width + 2), 55)
    ws.freeze_panes = ws.cell(row=r0 + 1, column=c0)


# Synthèse
ws = wb.active; ws.title = "Synthèse"
ws["A1"] = "Modèles de probabilité de défaut PME France - horizons 1 à 5 ans"
ws["A1"].font = Font(bold=True, size=14, color="1F3A5F")
lines = [
    "Données financières : « Ratios Financiers (BCE / INPI) », DGE, data.economie.gouv.fr (Licence Ouverte), liasses C et S, clôtures 2016-2025.",
    "Données d'identification : SIRENE (INSEE) — âge, forme juridique, secteur, région ; tranche d'effectif utilisée uniquement pour le périmètre.",
    "Défauts : BODACC (DILA), premier jugement d'ouverture de redressement, liquidation judiciaire ou sauvegarde.",
    "Date d'observation : clôture + 7 mois (délai légal d'approbation et de dépôt des comptes).",
    "Périmètre PME : sociétés commerciales, CA ≤ 50 M€, moins de 250 salariés, hors finance (K), administration (O), extraterritorial (U).",
    "Modèle champion : régression logistique sur variables en WoE, une par horizon ; challenger : LightGBM.",
    "PD cumulées rendues croissantes avec l'horizon (maximum cumulé) avant affectation à l'échelle de notation.",
    f"Échantillons : apprentissage {summ['train']:,} | test {summ['test']:,} | hors-période {summ['oot']:,} (clôtures ≥ {summ['oot_date']}).".replace(",", " "),
    f"Part des sociétés dont les PD brutes des 5 modèles ne sont pas croissantes : {summ['term_structure_violation_share']:.1%}.",
]
for i, t in enumerate(lines, start=3):
    ws.cell(row=i, column=1, value="• " + t)
ws.column_dimensions["A"].width = 130
key = perf[perf.echantillon.isin(["Logit - test", "Logit - OOT", "LightGBM - test"])] \
    .pivot(index="horizon", columns="echantillon", values="Gini").reset_index()
key.columns = ["Horizon", "Gini LightGBM test", "Gini Logit OOT", "Gini Logit test"]
key = key[["Horizon", "Gini Logit test", "Gini Logit OOT", "Gini LightGBM test"]]
ws.cell(row=14, column=1, value="Gini par horizon").font = Font(bold=True)
write_df(ws, key, r0=15, pct=key.columns[1:])
ws.column_dimensions["A"].width = 130
ws.freeze_panes = None
ws.add_image(XLImage(f"{OUT}/fig_roc.png"), "A23")

# Performance
ws = wb.create_sheet("Performance")
p = perf.copy()
write_df(ws, p, pct=["taux_observe", "pd_moyenne", "AUC", "Gini", "KS"], dec4=["Brier", "HL_pvalue"], dec2=["HL_stat"])

# Coefficients
ws = wb.create_sheet("Coefficients")
write_df(ws, coefs, dec4=["coef", "std_err", "p_value", "IV"], dec2=["VIF"])

# Grilles WoE (scorecards)
ws = wb.create_sheet("Grilles WoE")
write_df(ws, woe, pct=["taux_defaut"], dec4=["woe", "coef", "contribution_logit"])

# IV
ws = wb.create_sheet("Information Value")
ivd = iv.reset_index().rename(columns={"index": "Variable"})
ivd = ivd.rename(columns={ivd.columns[0]: "Variable"})
write_df(ws, ivd, dec4=list(ivd.columns[1:]))

# Calibration
ws = wb.create_sheet("Calibration")
write_df(ws, dec[["horizon", "decile", "n", "defauts", "pd_moyenne", "taux_observe"]],
         pct=["pd_moyenne", "taux_observe"])
ws.add_image(XLImage(f"{OUT}/fig_calibration.png"), "H2")

# Échelle de notation
ws = wb.create_sheet("Échelle de notation")
g = grades.copy()
write_df(ws, g, pct=[c for c in g.columns if c.startswith(("pd_", "dr_", "borne"))])
ws.add_image(XLImage(f"{OUT}/fig_term_structure.png"), "B16")

# LightGBM
ws = wb.create_sheet("Importance LightGBM")
im = imp.reset_index(); im = im.rename(columns={im.columns[0]: "Variable"})
write_df(ws, im, pct=list(im.columns[1:]))

# Comparaison avec la version 1 (compte de résultat seul, jeu Kaggle)
if len(sys.argv) > 2:
    v1 = pd.read_csv(sys.argv[2])
    comp = []
    for h in H:
        r1 = v1[(v1.horizon == f"{h}Y") & (v1.echantillon == "Logit - test")].iloc[0]
        r2 = perf[(perf.horizon == f"{h}Y") & (perf.echantillon == "Logit - test")].iloc[0]
        o2 = perf[(perf.horizon == f"{h}Y") & (perf.echantillon == "Logit - OOT")]
        comp.append({"Horizon": f"{h}Y", "Gini V1 (compte de résultat, Kaggle)": r1.Gini,
                     "Gini V2 (bilan, test)": r2.Gini,
                     "Gini V2 (bilan, hors-période)": o2.Gini.iloc[0] if len(o2) else np.nan,
                     "Gain (points)": (r2.Gini - r1.Gini) * 100})
    comp = pd.DataFrame(comp)
    ws = wb.create_sheet("Comparaison V1-V2", 1)
    write_df(ws, comp, pct=[c for c in comp.columns if c.startswith("Gini")], dec2=["Gain (points)"])
    comp.to_csv(f"{OUT}/comparaison_v1_v2.csv", index=False)

wb.save(f"{OUT}/PD_PME_France_modeles_bilan.xlsx")
print("ok")
