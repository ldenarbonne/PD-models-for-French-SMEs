"""
14_qualite_donnees.py — Onglet « Qualité des données » : taux de manquants et redondances
(échantillon aléatoire de 10 % de la base de modélisation).
"""
import sys
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill

BASE, XLSX = sys.argv[1], sys.argv[2]
smp = pd.concat([c.sample(frac=0.1, random_state=1)
                 for c in pd.read_csv(BASE, chunksize=500_000, low_memory=False)])
cols = [c for c in smp.columns if c not in ("siren", "cloture", "t0", "date_defaut", "eff_tranche")
        and not c.startswith("def_")]
conf = smp["comptes_confidentiels"] == 1
miss = pd.DataFrame({"Manquants (total)": smp[cols].isna().mean(),
                     "Manquants hors comptes confidentiels": smp.loc[~conf, cols].isna().mean()}) \
    .sort_values("Manquants (total)", ascending=False)
num = [c for c in smp[cols].select_dtypes("number").columns]
cor = smp[num].corr(method="spearman")
pairs = pd.DataFrame([(a, b, cor.loc[a, b]) for i, a in enumerate(num) for b in num[i + 1:]
                      if abs(cor.loc[a, b]) >= 0.5], columns=["Variable 1", "Variable 2", "Rho de Spearman"]) \
    .sort_values("Rho de Spearman", key=abs, ascending=False)

wb = load_workbook(XLSX)
ws = wb.create_sheet("Qualité des données", 2)
hdr, hf = PatternFill("solid", fgColor="1F3A5F"), Font(bold=True, color="FFFFFF")
ws["A1"] = f"Échantillon aléatoire de 10 % : {len(smp):,} exercices".replace(",", " ")
ws["A1"].font = Font(bold=True)
for j, h in enumerate(["Variable"] + list(miss.columns), start=1):
    c = ws.cell(row=3, column=j, value=h); c.fill, c.font = hdr, hf
for i, (v, r) in enumerate(miss.iterrows(), start=4):
    ws.cell(row=i, column=1, value=v)
    for j, x in enumerate(r, start=2):
        ws.cell(row=i, column=j, value=float(x)).number_format = "0.0%"
c0 = 6
for j, h in enumerate(pairs.columns, start=c0):
    c = ws.cell(row=3, column=j, value=h); c.fill, c.font = hdr, hf
for i, r in enumerate(pairs.itertuples(index=False), start=4):
    ws.cell(row=i, column=c0, value=r[0]); ws.cell(row=i, column=c0 + 1, value=r[1])
    ws.cell(row=i, column=c0 + 2, value=float(r[2])).number_format = "+0.00"
for col, w in zip("ABCDEFGH", [26, 18, 30, 4, 4, 24, 24, 16]):
    ws.column_dimensions[col].width = w
wb.save(XLSX)
print(miss.round(3).head(), pairs.head(), sep="\n")
